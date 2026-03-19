#!/usr/bin/env python3
"""
自动化配平薪酬：基于“序时账”与“应付职工薪酬分配检查情况表”自动生成
Excel 第三张表：配平薪酬。实现要点：
- 统计每月序时账中的薪酬相关借方总额
- 将应付表中的账面计提金额按月分摊，得到每月应付分配额（按月占比法）
- 计算差异，并给出差异原因汇总与默认调整科目
- 写入一个新工作表：配平薪酬，字段包括：期间、序时账薪酬总额、应付分配额、差异、差异原因、调整科目、备注
"""
from __future__ import annotations

import math
from collections import defaultdict
from pathlib import Path

import openpyxl


KW = [
    "职工薪酬",
    "工资",
    "奖金",
    "津贴",
    "补贴",
    "福利",
]

DEFAULT_ADJ = "66020101 管理费用_职工薪酬_工资及奖金"
SHEET_PAYROLL_JRN = "序时账"
SHEET_APAY = "应付职工薪酬分配检查情况表 "
SHEET_OUT = "配平薪酬"


def is_payroll_subject(text: str | None) -> bool:
    if not text:
        return False
    t = text
    t = t if isinstance(t, str) else str(t)
    return any(k in t for k in KW)


def parse_number(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(v)
    except Exception:
        # sometimes numbers are strings with commas
        try:
            return float(str(v).replace(',', ''))
        except Exception:
            return 0.0


def main():
    path = Path.cwd() / "薪酬勾稽表.xlsx"
    if not path.exists():
        print(f"Excel 文件未找到: {path}")
        return

    wb = openpyxl.load_workbook(path, data_only=True)

    # 1) 序时账：聚合每月薪酬相关借方总额
    if SHEET_PAYROLL_JRN not in wb.sheetnames:
        print(f"找不到工作表: {SHEET_PAYROLL_JRN}")
        return
    jrn = wb[SHEET_PAYROLL_JRN]
    # detect header indices
    header = [c.value for c in next(jrn.iter_rows(min_row=1, max_row=1))]
    idx_period = next((i for i, v in enumerate(header) if str(v).strip() == "期间"), None)
    idx_dr = next((i for i, v in enumerate(header) if str(v).strip() == "借方本币"), None)
    idx_dst = next((i for i, v in enumerate(header) if str(v).strip() == "对方科目名称"), None)
    if idx_period is None or idx_dr is None or idx_dst is None:
        print("序时账表头未发现关键列（期间/借方本币/对方科目名称），无法计算。")
        return

    payroll_by_period = defaultdict(float)
    for row in jrn.iter_rows(min_row=2, values_only=True):
        period = row[idx_period]
        dst = row[idx_dst]
        amt = row[idx_dr]
        if period is None:
            continue
        if is_payroll_subject(str(dst) if dst is not None else None):
            payroll_by_period[str(period)] += parse_number(amt)

    # 2) 应付表：提取工资/奖金/津贴等科目的账面计提金额，累计总额
    if SHEET_APAY not in wb.sheetnames:
        print(f"找不到工作表: {SHEET_APAY}")
        return
    apay = wb[SHEET_APAY]
    # locate header row for 该表格的列名
    ap_header = [c.value for c in next(apay.iter_rows(min_row=1, max_row=1))]
    idx_item = next((i for i, v in enumerate(ap_header) if str(v).strip() == "项  目"), None)
    idx_cost_base = next((i for i, v in enumerate(ap_header) if str(v).strip() == "账面计提金额"), None)
    idx_reason = next((i for i, v in enumerate(ap_header) if str(v).strip() == "差异原因"), None)
    # 仅聚合 payroll-related 项目
    acct_cols_start = idx_cost_base
    payroll_rows = []
    per_item_values = []  # (name, value, reason)
    for r in apay.iter_rows(min_row=2, values_only=True):
        item = r[idx_item] if idx_item is not None else None
        if item is None:
            continue
        if isinstance(item, str) and item.strip().startswith("工资"):
            val = r[idx_cost_base] if idx_cost_base is not None else None
            reason = r[idx_reason] if idx_reason is not None else None
            payroll_rows.append((item, parse_number(val), reason))
            per_item_values.append((item, parse_number(val), reason))
        # 其它相关科目也可能包含在工资相关分类中，保守地收集所有以"工资|奖金|福利|社保|公积金|补贴|津贴"等打头的项
        elif isinstance(item, str) and any(item.strip().startswith(w) for w in ["工资", "奖金", "福利", "社保", "公积金", "津贴", "补贴"]):
            val = r[idx_cost_base] if idx_cost_base is not None else None
            reason = r[idx_reason] if idx_reason is not None else None
            payroll_rows.append((item, parse_number(val), reason))
            per_item_values.append((item, parse_number(val), reason))

    total_acct_payroll = sum(val for _, val, _ in payroll_rows)

    # 3) 计算按月分摊（按月薪酬总额占比）
    months = sorted([m for m in payroll_by_period.keys() if m is not None])
    total_payroll_all = sum(payroll_by_period[m] for m in months) if months else 0.0

    allocations = {}
    if total_payroll_all > 0:
        for m in months:
            allocations[m] = total_acct_payroll * (payroll_by_period[m] / total_payroll_all)
    # 4) 写入新表：配平薪酬
    if SHEET_OUT in wb.sheetnames:
        ws = wb[SHEET_OUT]
    else:
        ws = wb.create_sheet(title=SHEET_OUT)
    # 表头
    headers = [
        "期间",
        "序时账薪酬总额",
        "应付分配额",
        "差异",
        "差异原因",
        "调整科目",
        "备注",
    ]
    ws.append(headers)
    # 统计差异原因：把应付表中所有非空的差异原因整合为一个字符串
    reasons = [r for _, _, r in payroll_rows if r]
    unique_reasons = sorted(set(str(x) for x in reasons if x))
    reason_text = "；".join(unique_reasons) if unique_reasons else ""
    for m in months:
        sched = payroll_by_period[m]
        alloc = allocations.get(m, 0.0)
        diff = sched - alloc
        row = [
            m,
            round(sched, 2) if not math.isinf(sched) else sched,
            round(alloc, 2) if not math.isinf(alloc) else alloc,
            round(diff, 2),
            reason_text,
            DEFAULT_ADJ,
            "自动生成，请核对凭证编号与科目归集",
        ]
        ws.append(row)

    wb.save(path)
    print(f"已生成/更新配平薪酬表：{path}，涉及月份: {', '.join(months)}")


if __name__ == "__main__":
    main()
